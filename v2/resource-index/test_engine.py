import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

spec=importlib.util.spec_from_file_location('resource_engine',Path(__file__).with_name('engine.py'))
e=importlib.util.module_from_spec(spec);spec.loader.exec_module(e)


class ResourceTests(unittest.TestCase):
    def setUp(self):
        (e.PROJECT/'tmp').mkdir(exist_ok=True)
        self.tmp=tempfile.TemporaryDirectory(dir=e.PROJECT/'tmp')
        self.root=Path(self.tmp.name)/'originals';self.root.mkdir();self.work=Path(self.tmp.name)/'candidate'

    def tearDown(self):self.tmp.cleanup()

    def seed(self):
        (self.root/'A3.txt').write_text('A3 产品说明书\n故障现象：升级后黑屏；检查电源。',encoding='utf8')
        (self.root/'A3L.txt').write_text('A3L 产品说明书\n故障现象：升级后黑屏；检查排线。',encoding='utf8')
        (self.root/'copy').mkdir();(self.root/'copy/A3.txt').write_bytes((self.root/'A3.txt').read_bytes())
        e.scan(self.root,self.work);return e.index_documents(self.work)

    def test_model_suffix_and_short_chinese(self):
        self.seed()
        r=e.query(self.work,'黑屏',mode='text',model='A3')
        self.assertTrue(r['results']);self.assertTrue(all('A3L' not in h['text'] for h in r['results']))
        self.assertTrue(e.query(self.work,'黑屏',mode='text',model='A3L')['results'])
        self.assertEqual(e.query(self.work,'黑屏',mode='text',model='VP620A')['results'],[])

    def test_identical_files_parse_once_and_keep_locations(self):
        r=self.seed();self.assertEqual(r['parsed'],2)
        h=e.query(self.work,'黑屏',mode='text',model='A3')['results'][0]
        self.assertEqual(len(h['locations']),2);self.assertEqual(e.report(self.work)['duplicateExtraFiles'],1)
        self.assertEqual(e.index_documents(self.work),{'cacheHits':2})

    def test_rename_relocates_identity_and_reuses_content(self):
        self.seed();db=e.connect(self.work,True)
        old=db.execute('SELECT resource_id,version_id FROM locations WHERE path=?',(str(self.root/'A3L.txt'),)).fetchone();db.close()
        (self.root/'A3L.txt').rename(self.root/'moved-A3L.txt');e.scan(self.root,self.work)
        db=e.connect(self.work,True);new=db.execute('SELECT resource_id,version_id FROM locations WHERE path=?',(str(self.root/'moved-A3L.txt'),)).fetchone()
        self.assertEqual(tuple(old),tuple(new));self.assertEqual(db.execute("SELECT state FROM locations WHERE name='A3L.txt'").fetchone()[0],'missing');db.close()
        self.assertEqual(e.index_documents(self.work),{'cacheHits':2})

    def test_evidence_round_trip_and_different_causes_not_merged(self):
        self.seed();a=e.query(self.work,'黑屏',mode='text',model='A3')['results'][0];b=e.query(self.work,'黑屏',mode='text',model='A3L')['results'][0]
        self.assertNotEqual(a['sha256'],b['sha256']);r=e.read_block(self.work,a['blockId'],0)
        self.assertEqual(r['blocks'][0]['text'],a['text']);self.assertEqual(r['blocks'][0]['locator'],a['locator'])

    def test_original_guard_and_no_binary_hash_or_cloud_invention(self):
        (self.root/'R712.bin').write_bytes(b'not a document');(self.root/'key.pem').write_text('PRIVATE');(self.root/'x.pdf.part').write_bytes(b'half')
        e.scan(self.root,self.work);db=e.connect(self.work,True)
        self.assertIsNone(db.execute("SELECT sha256 FROM locations WHERE name='R712.bin'").fetchone()[0]);self.assertEqual(db.execute('SELECT count(*) FROM locations WHERE cloud_id IS NOT NULL').fetchone()[0],0);db.close()
        with self.assertRaises(ValueError):e.scan(self.root,self.root/'bad')
        with self.assertRaises(ValueError):e.connect(self.root.parent.parent.parent/'outside')
        self.assertEqual(e.query(self.work,'PRIVATE')['results'],[])

    def test_docx_table_and_pptx_locators(self):
        from docx import Document
        from pptx import Presentation
        d=Document();d.add_heading('技术条件',1);t=d.add_table(rows=2,cols=2);t.cell(0,0).text='型号';t.cell(0,1).text='限制';t.cell(1,0).text='VP620A';t.cell(1,1).text='必须核对硬件版本';p=self.root/'test.docx';d.save(p)
        rows=e.extract(p)['blocks'];self.assertTrue(any(b['locator'].get('row')==2 and 'VP620A' in b['text'] for b in rows))
        ppt=Presentation();s=ppt.slides.add_slide(ppt.slide_layouts[1]);s.shapes.title.text='培训';s.placeholders[1].text='升级前检查版本';p=self.root/'test.pptx';ppt.save(p)
        self.assertTrue(any(b['locator']['slide']==1 and '升级前' in b['text'] for b in e.extract(p)['blocks']))

    def test_xlsx_cells_and_csv_quotes(self):
        from openpyxl import Workbook
        wb=Workbook();wb.active.title='版本条件';wb.active.append(['型号','条件']);wb.active.append(['VP620A','V1.2']);p=self.root/'v.xlsx';wb.save(p);wb.close()
        blocks=e.extract(p)['blocks'];self.assertEqual(blocks[1]['locator']['sheet'],'版本条件');self.assertEqual(blocks[1]['locator']['cells'][1]['cell'],'B2')
        p=self.root/'cases.csv';p.write_text('现象,处理\n"黑屏,偶发",检查供电\n',encoding='utf8');r=e.extract(p)
        self.assertEqual(r['blocks'][1]['locator']['cells'][0],'黑屏,偶发')

    def test_changed_source_refuses_stale_hash_and_records_failure(self):
        (self.root/'A3.txt').write_text('原始资料',encoding='utf8');e.scan(self.root,self.work)
        (self.root/'A3.txt').write_text('已经变化的资料',encoding='utf8')
        self.assertEqual(e.index_documents(self.work)['failed'],1)

    def test_exported_xlsx_wrong_dimensions_do_not_hide_rows(self):
        from openpyxl import Workbook
        wb=Workbook();wb.active['E1']='固件版本';wb.active['E4']='V23.06.01.30';p=self.root/'export.xlsx';wb.save(p);wb.close()
        with zipfile.ZipFile(p) as z:parts={n:z.read(n) for n in z.namelist()}
        parts['xl/worksheets/sheet1.xml']=parts['xl/worksheets/sheet1.xml'].replace(b'E1:E4',b'A1:A1')
        with zipfile.ZipFile(p,'w') as z:
            for n,b in parts.items():z.writestr(n,b)
        rows=e.extract(p)['blocks'];self.assertEqual(len(rows),2);self.assertEqual(rows[1]['locator']['row'],4);self.assertIn('V23.06.01.30',rows[1]['text'])

    def test_spreadsheet_coordinates_and_foreign_mentions_are_not_model_identity(self):
        from openpyxl import Workbook
        wb=Workbook();wb.active['A3']='固件升级';p=self.root/'firmware.xlsx';wb.save(p);wb.close()
        (self.root/'VP620A.txt').write_text('VP620A参数\n配合A3使用的说明',encoding='utf8')
        e.scan(self.root,self.work);e.index_documents(self.work)
        self.assertFalse(e.query(self.work,'A3',mode='files')['results'])
        self.assertNotIn('A3',e.model_tokens(e.extract(p)['blocks'][0]['text']))

    def test_rescan_and_new_duplicate_preserve_cached_body_model_links(self):
        (self.root/'generic.txt').write_text('R712 故障现象：黑屏',encoding='utf8')
        e.scan(self.root,self.work);e.index_documents(self.work);e.scan(self.root,self.work)
        self.assertTrue(e.query(self.work,'黑屏',mode='text',model='R712')['results'])
        (self.root/'copy.txt').write_bytes((self.root/'generic.txt').read_bytes());e.scan(self.root,self.work)
        r=e.query(self.work,'黑屏',mode='text',model='R712')['results'];self.assertEqual(len(r[0]['locations']),2)

    def test_filename_rank_precedes_candidate_limit(self):
        for i in range(1050):(self.root/('产品说明资料%04d.txt'%i)).write_text('普通资料',encoding='utf8')
        (self.root/'ZZ软件操作说明资料.txt').write_text('目标原文',encoding='utf8')
        e.scan(self.root,self.work)
        result=e.query(self.work,'ZZ软件操作说明资料',mode='files')
        self.assertEqual(result['results'][0]['name'],'ZZ软件操作说明资料.txt')

    def test_office_lock_and_misnamed_legacy_are_explicit_pending(self):
        (self.root/'~$lock.docx').write_bytes(b'office lock')
        p=self.root/'old.docx';p.write_bytes(bytes.fromhex('d0cf11e0a1b11ae1')+b'legacy')
        e.scan(self.root,self.work);db=e.connect(self.work,True)
        self.assertEqual(db.execute("SELECT state FROM locations WHERE name='~$lock.docx'").fetchone()[0],'office-lock-file');db.close()
        self.assertEqual(e.extract(p)['state'],'misnamed-legacy-office')

    def test_exact_version_not_prefix_and_standard_hd_alias(self):
        for v in ['V1.0','V1.01']:(self.root/('A3 '+v+'.txt')).write_text('A3参数',encoding='utf8')
        e.scan(self.root,self.work)
        rows=e.query(self.work,'A3',mode='files',model=' HD_A3 ',version='V1.0')['results']
        self.assertEqual([r['name'] for r in rows],['A3 V1.0.txt'])


if __name__=='__main__':unittest.main()
